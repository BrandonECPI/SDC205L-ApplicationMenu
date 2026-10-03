from datetime import datetime
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference

menu_options = ["Input Data", "View Current Data", "Generate Report"]


def convertData(value):
    # convertData takes the temperature in Fahrenheit and returns it converted to Celsius
    return (value - 32) * 5 / 9


def insertData(path, data):
    # insertData takes a file path and a comma-separated string of data, opens the file with
    # append-only permission (creating it if it does not already exist), and writes the data to it
    try:
        with open(path, "a") as file:
            file.write(data + "\n")
    except Exception as e:
        print("Error writing to file: " + str(e))
        raise


def viewData(path):
    # viewData takes a file path, opens it with read-only permission, and displays the file's
    # path along with its contents
    try:
        with open(path, "r") as file:
            print("The file " + path)
            for line in file:
                print(line.strip())
    except Exception as e:
        print("Error reading file: " + str(e))


def getInput():
    # getInput asks how many entries to add, then for each one collects a date and temperature,
    # converts it, and saves it to ZooData.csv using insertData
    num_entries = int(input("How many entries are you inputting? "))
    for i in range(num_entries):
        date = input("Enter a date: ")
        temp = int(input("Enter the highest temp for the inputted date: "))
        # convertData takes one argument, the temperature in Fahrenheit, and returns the temperature converted to Celsius
        converted = convertData(temp)
        data = date + "," + str(temp) + "," + str(converted)
        try:
            insertData("ZooData.csv", data)
            print("The following was saved at " + str(datetime.now()) + " : ")
            print(data)
        except Exception:
            pass


def createChart(path, chart_type):
    # createChart takes the path to the csv data file (str) and a chart type of "line" or "bar" (str).
    # It reads the data, saves it to final.xlsx using openpyxl, and adds the requested chart. No return value.
    source = input("Enter 'F' to graph the Fahrenheit data or 'C' to graph the Celsius data: ").strip().upper()

    dates = []
    values = []

    with open(path, "r") as file:
        for line in file:
            parts = line.strip().split(",")
            date = parts[0]
            tempF = float(parts[1])
            tempC = float(parts[2])
            dates.append(date)
            if source == "F":
                values.append(tempF)
            else:
                values.append(tempC)

    wb = Workbook()
    ws = wb.active
    ws.append(["Date", "Fahrenheit" if source == "F" else "Celsius"])
    for d, v in zip(dates, values):
        ws.append([d, v])

    if chart_type == "bar":
        chart = BarChart()
    else:
        chart = LineChart()

    chart.title = "bralaz4151 " + datetime.now().strftime("%m/%d/%Y")
    chart.x_axis.title = "Date"
    chart.y_axis.title = "Temperature (F)" if source == "F" else "Temperature (C)"

    data_ref = Reference(ws, min_col=2, min_row=1, max_row=ws.max_row)
    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)

    ws.add_chart(chart, "D2")

    wb.save("final.xlsx")


def generateReport(path):
    # generateReport takes the path to the csv data file (str), asks the user to choose a chart
    # type, and calls createChart with that path and type. No return value.
    chart_type = input("Which graph would you like to create? Enter 'line' or 'bar': ").strip().lower()
    createChart(path, chart_type)


print("bralaz4151 Spreadsheet Automation Menu")
print("Choose a number from the following options")

# option represents each menu item as the loop prints it along with its number
for index, option in enumerate(menu_options, start=1):
    print(str(index) + ". " + option)

# The next line retrieves the inputted option and stores into the variable called choice.
choice = input()

if choice.isdigit() and 1 <= int(choice) <= len(menu_options):
    print("You selected " + choice + " at " + str(datetime.now()))
    if choice == "1":
        getInput()
    elif choice == "2":
        viewData("ZooData.csv")
    elif choice == "3":
        generateReport("ZooData.csv")
else:
    print("Error: Invalid choice selected.")